"""Tests for the generated-document writers (PDF, Word, PowerPoint, Excel, text),
the shared Markdown parser behind them, and the copy-card prompts."""
import csv
import io
import re
import unicodedata
import zipfile
from datetime import datetime, timezone
from unittest.mock import Mock, patch
from types import SimpleNamespace

from django.core.cache import cache
from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from myapp import ai_chat, doc_blocks, doc_pdf, doc_sheet, doc_slides, doc_text, doc_word, file_convert
from myapp.models import AIGeneratedFile
from myapp.views import (
    _ai_excel_bytes, _ai_pdf_bytes, _ai_powerpoint_bytes, _ai_word_document_bytes,
    _extract_ai_generated_file_content, _ai_generated_file_instruction,
)

MADE_ON = datetime(2026, 10, 6, 9, 30, tzinfo=timezone.utc)

REPORT = """# Digital Marketing Plan

A practical 90-day plan for **Sweet Crumbs Bakery**.

> **Note:** All budgets are in ₹ and exclude GST.

## 1. Goals

- Reach **300** online orders
  - Reply within 2 hours
- Collect 200 reviews

## 2. Budget

| Channel | Budget (₹) | Leads |
|:--|--:|--:|
| Meta Ads | 15,000 | 300 |
| SEO | 5,000 | 60 |
| **Total** | **20,000** | **360** |

```chart
type: column
title: Leads by channel
Channel, Leads
Meta Ads, 300
SEO, 60
```

## 3. Steps

1. Claim the profile
2. Add photos

## 4. Sources

- [Meta Help](https://www.facebook.com/business/help)
"""


def pdf_text(data):
    """All the text of a PDF. The layout engine sets "fi" and "ff" as single
    ligature glyphs, which extract as one character (ﬁ), so they are folded
    back to plain letters before comparing."""
    raw = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages)
    return unicodedata.normalize('NFKC', raw)


class MarkdownParserTests(SimpleTestCase):
    def test_blocks_for_a_typical_model_document(self):
        blocks = doc_blocks.parse_markdown(REPORT)
        kinds = [block.kind for block in blocks]
        self.assertEqual(kinds[0], 'heading')
        self.assertEqual(doc_blocks.document_title(blocks), 'Digital Marketing Plan')
        self.assertIn('table', kinds)
        self.assertIn('chart', kinds)
        callout = next(block for block in blocks if block.kind == 'quote')
        self.assertEqual(callout.callout, 'note')
        self.assertTrue(callout.text.startswith('All budgets'))

    def test_table_alignment_and_numeric_columns(self):
        table = next(b for b in doc_blocks.parse_markdown(REPORT) if b.kind == 'table')
        self.assertEqual(table.header, ['Channel', 'Budget (₹)', 'Leads'])
        # Separator colons are honoured; a column of numbers aligns right on its own.
        self.assertEqual(doc_blocks.column_aligns(table), ['left', 'right', 'right'])

    def test_a_centre_separator_with_a_single_dash_is_still_a_table(self):
        blocks = doc_blocks.parse_markdown('| A | B |\n|:-:|:-:|\n| 1 | 2 |')
        self.assertEqual([b.kind for b in blocks], ['table'])
        self.assertEqual(blocks[0].aligns, ['center', 'center'])

    def test_numbered_lists_restart_and_nest(self):
        items = [b for b in doc_blocks.parse_markdown('1. a\n2. b\n\ntext\n\n1. c\n   - sub') if b.kind == 'item']
        self.assertEqual([(i.ordered, i.number, i.level) for i in items], [(True, 1, 0), (True, 2, 0), (True, 1, 0), (False, 0, 1)])

    def test_inline_formatting_and_escapes(self):
        spans = doc_blocks.inline_spans(r'Hi **bold *both*** `code` [site](https://a.example/x) \*literal\*')
        by_text = {span.text: span for span in spans}
        self.assertTrue(by_text['bold '].bold)
        self.assertTrue(by_text['both'].bold and by_text['both'].italic)
        self.assertTrue(by_text['code'].code)
        self.assertEqual(by_text['site'].url, 'https://a.example/x')
        self.assertIn('*literal*', ''.join(span.text for span in spans))

    def test_raw_data_survives_escape_markdown_unchanged(self):
        raw = '**NOTE** a_b_c [x](y) <b>t</b> 5*3 | pipe हिन्दी ₹5'
        self.assertEqual(doc_blocks.plain_text(doc_blocks.escape_markdown(raw)), raw)

    def test_numbers(self):
        self.assertEqual(doc_blocks.parse_number('₹1,20,000'), 120000.0)
        self.assertEqual(doc_blocks.parse_number('12.5%'), 12.5)
        self.assertEqual(doc_blocks.parse_number('(45)'), -45.0)
        self.assertIsNone(doc_blocks.parse_number('n/a'))

    def test_chart_spec(self):
        spec = doc_blocks.parse_chart('type: pie\ntitle: Share\nChannel, Share\nA, 60\nB, 40')
        self.assertEqual(spec['type'], 'pie')
        self.assertEqual(spec['labels'], ['A', 'B'])
        self.assertEqual(spec['series'][0]['values'], [60.0, 40.0])
        self.assertIsNone(doc_blocks.parse_chart('just words'))

    def test_control_characters_are_dropped(self):
        blocks = doc_blocks.parse_markdown('# Title\x00\n\nBody\x0c text')
        self.assertNotIn('\x00', blocks[0].text)
        self.assertNotIn('\x0c', blocks[1].text)


class PdfWriterTests(SimpleTestCase):
    def test_designed_pdf_carries_the_content(self):
        data = doc_pdf.render_pdf(REPORT, brand='Vidhyora', created=MADE_ON)
        text = pdf_text(data)
        self.assertTrue(data.startswith(b'%PDF'))
        for expected in ('Digital Marketing Plan', 'Sweet Crumbs Bakery', 'Meta Ads', '15,000', 'Claim the profile', 'Meta Help'):
            self.assertIn(expected, text)
        self.assertIn('Prepared on 6 October 2026', text)
        self.assertIn('Page 1 of', text)
        for marker in ('**', '```', '](', '|---'):
            self.assertNotIn(marker, text)

    def test_links_are_clickable_and_the_domain_is_shown(self):
        import pymupdf
        data = doc_pdf.render_pdf(REPORT, created=MADE_ON)
        document = pymupdf.open('pdf', data)
        uris = [link.get('uri') for page in document for link in page.get_links()]
        self.assertIn('https://www.facebook.com/business/help', uris)
        self.assertIn('facebook.com', pdf_text(data))

    def test_headings_become_bookmarks(self):
        import pymupdf
        outline = pymupdf.open('pdf', doc_pdf.render_pdf(REPORT, created=MADE_ON)).get_toc()
        titles = [entry[1] for entry in outline]
        self.assertEqual(titles[0], 'Digital Marketing Plan')
        self.assertIn('2. Budget', titles)

    def test_a_long_table_repeats_its_header_on_every_page(self):
        rows = '\n'.join(f'| Item {n} | {n * 3} |' for n in range(1, 90))
        data = doc_pdf.render_pdf(f'# Inventory\n\n| Product | Units |\n|---|---:|\n{rows}\n')
        reader = PdfReader(io.BytesIO(data))
        self.assertGreater(len(reader.pages), 1)
        for page in reader.pages:
            self.assertIn('Product', page.extract_text())
        # Every row is there exactly once and none is cut in half.
        text = pdf_text(data)
        for n in (1, 45, 89):
            self.assertEqual(len(re.findall(rf'Item {n}\b', text)), 1)

    def test_a_long_document_gets_a_contents_list(self):
        body = '\n\n'.join(f'## Section {n}\n\n' + 'Lorem ipsum dolor sit amet. ' * 120 for n in range(1, 7))
        data = doc_pdf.render_pdf(f'# Handbook\n\n{body}')
        self.assertGreaterEqual(len(PdfReader(io.BytesIO(data)).pages), 3)
        self.assertIn('Contents', pdf_text(data))

    def test_unbranded_output_has_no_title_band_or_brand(self):
        text = pdf_text(doc_pdf.render_pdf('# Invoice\n\nTotal 4999.', brand='Vidhyora', branded=False))
        self.assertIn('Invoice', text)
        self.assertNotIn('Vidhyora', text)
        self.assertNotIn('Prepared on', text)

    def test_empty_content_is_a_blank_page(self):
        self.assertEqual(pdf_text(doc_pdf.render_pdf('')).strip(), '')

    def test_hindi_text_is_embedded_not_dropped(self):
        import pymupdf
        data = doc_pdf.render_pdf('# सारांश\n\nयह योजना ₹30,000 की है।')
        text = pymupdf.open('pdf', data)[0].get_text()
        self.assertIn('सारांश', text)
        self.assertIn('₹30,000', text)

    def test_fonts_are_subset_so_the_file_stays_small(self):
        self.assertLess(len(doc_pdf.render_pdf(REPORT, created=MADE_ON)), 400_000)

    def test_a_long_unbroken_string_wraps_instead_of_running_off_the_page(self):
        import pymupdf
        page = pymupdf.open('pdf', doc_pdf.render_pdf(
            '# T\n\n| A | B |\n|---|---|\n| ' + 'x' * 400 + ' | visible |\n'))[0]
        words = page.get_text('words')
        self.assertLessEqual(max(word[2] for word in words), 545)      # the right margin is 541
        self.assertIn('visible', [word[4] for word in words])           # the next column is still on the page


class WordWriterTests(SimpleTestCase):
    def setUp(self):
        self.data = doc_word.render_docx(REPORT, brand='Vidhyora', created=MADE_ON)
        self.document = Document(io.BytesIO(self.data))
        self.xml = zipfile.ZipFile(io.BytesIO(self.data)).read('word/document.xml').decode('utf-8')

    def test_real_styles_and_text(self):
        styles = {p.style.name for p in self.document.paragraphs}
        for expected in ('Heading 1', 'Heading 2', 'List Bullet', 'List Bullet 2', 'List Number'):
            self.assertIn(expected, styles)
        texts = [p.text for p in self.document.paragraphs]
        self.assertIn('Digital Marketing Plan', texts)
        self.assertTrue(any(t.startswith('Prepared on 6 October 2026') for t in texts))
        self.assertFalse(any('**' in t for t in texts))

    def test_table_has_a_repeating_header_and_shaded_cells(self):
        table = self.document.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells], ['Channel', 'Budget (₹)', 'Leads'])
        self.assertIn('<w:tblHeader/>', self.xml)
        self.assertIn('w:fill="059669"', self.xml)           # header
        self.assertIn('w:fill="f3f8f6"', self.xml)           # zebra row
        self.assertIn('w:val="right"', self.xml)             # numbers align right

    def test_chart_is_an_embedded_picture_with_alt_text(self):
        self.assertIn('descr="Leads by channel"', self.xml)
        self.assertEqual(len(self.document.inline_shapes), 1)

    def test_links_are_real_hyperlinks(self):
        self.assertIn('<w:hyperlink', self.xml)
        relationships = zipfile.ZipFile(io.BytesIO(self.data)).read('word/_rels/document.xml.rels').decode()
        self.assertIn('https://www.facebook.com/business/help', relationships)

    def test_numbered_lists_restart(self):
        data = doc_word.render_docx('1. a\n2. b\n\nbreak\n\n1. c\n2. d')
        ids = set(re.findall(r'<w:numId w:val="(\d+)"/>', zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml').decode()))
        self.assertEqual(len(ids), 2)                         # two lists, two counters

    def test_footer_has_live_page_numbers_and_properties(self):
        footer = zipfile.ZipFile(io.BytesIO(self.data)).read('word/footer1.xml').decode()
        self.assertIn('PAGE', footer)
        self.assertIn('NUMPAGES', footer)
        self.assertEqual(self.document.core_properties.title, 'Digital Marketing Plan')
        self.assertEqual(self.document.core_properties.author, 'Vidhyora')

    def test_unbranded_conversion_has_no_brand_or_title_block(self):
        document = Document(io.BytesIO(doc_word.render_docx('# Invoice\n\nTotal 4999.', brand='Vidhyora', branded=False)))
        texts = [p.text for p in document.paragraphs]
        self.assertIn('Invoice', texts)
        self.assertFalse(any('Prepared on' in t or 'VIDHYORA' in t for t in texts))

    def test_properties_are_in_schema_order(self):
        # Word rejects a paragraph whose properties are out of order: shading must
        # follow borders, which must follow keepLines.
        paragraph = re.search(r'<w:pPr><w:keepLines/>(.*?)</w:pPr>', self.xml).group(1)
        self.assertLess(paragraph.index('<w:pBdr>'), paragraph.index('<w:shd '))


class SlidesWriterTests(SimpleTestCase):
    DECK = (
        '# Growing Online\n- A 90-day plan\n\n# Today\n- 120 orders a month\n  - mostly word of mouth\n'
        'Notes: ask about walk-ins\n\n# Budget\n| Channel | Spend |\n|---|--:|\n| Meta | 15,000 |\n| SEO | 5,000 |\n\n'
        '# Leads\n```chart\ntype: column\ntitle: Leads\nChannel, Leads\nMeta, 300\nSEO, 60\n```\n'
    )

    def setUp(self):
        self.deck = Presentation(io.BytesIO(doc_slides.render_pptx(self.DECK, brand='Vidhyora', created=MADE_ON)))

    def test_slides_titles_and_layouts(self):
        self.assertEqual([s.shapes.title.text for s in self.deck.slides], ['Growing Online', 'Today', 'Budget', 'Leads'])
        self.assertEqual(self.deck.slides[0].slide_layout.name, 'Title Slide')
        self.assertEqual(self.deck.slides[1].slide_layout.name, 'Title and Content')
        self.assertEqual(self.deck.slide_width, int(13.333 * 914400))

    def test_bullets_levels_and_notes(self):
        paragraphs = self.deck.slides[1].placeholders[1].text_frame.paragraphs
        self.assertEqual([p.text for p in paragraphs], ['120 orders a month', 'mostly word of mouth'])
        self.assertEqual(self.deck.slides[1].notes_slide.notes_text_frame.text, 'ask about walk-ins')

    def test_table_slide_has_a_real_table(self):
        tables = [s for s in self.deck.slides[2].shapes if getattr(s, 'has_table', False) and s.has_table]
        self.assertEqual(len(tables), 1)
        table = tables[0].table
        self.assertEqual([c.text for c in table.rows[0].cells], ['Channel', 'Spend'])
        self.assertEqual(table.cell(1, 1).text, '15,000')

    def test_chart_slide_has_a_native_chart(self):
        charts = [s for s in self.deck.slides[3].shapes if getattr(s, 'has_chart', False) and s.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertEqual(list(charts[0].chart.plots[0].categories), ['Meta', 'SEO'])

    def test_every_shape_stays_on_the_slide(self):
        for slide in self.deck.slides:
            for shape in slide.shapes:
                if shape.left is None:
                    continue   # inherits its place from the layout
                self.assertGreaterEqual(shape.left, 0, shape.name)
                self.assertLessEqual(shape.left + shape.width, self.deck.slide_width + 10, shape.name)
                self.assertLessEqual(shape.top + shape.height, self.deck.slide_height + 10, shape.name)

    def test_too_many_points_continue_on_another_slide(self):
        points = '\n'.join(f'- Point number {n} explains something in a fairly long sentence for the audience' for n in range(1, 16))
        deck = Presentation(io.BytesIO(doc_slides.render_pptx(f'# Cover\n- Sub\n\n# Many\n{points}\n')))
        titles = [s.shapes.title.text for s in deck.slides]
        self.assertEqual(titles[1], 'Many')
        self.assertTrue(all(t.endswith('(cont.)') for t in titles[2:]) and len(titles) > 2)
        shown = sum(len(s.placeholders[1].text_frame.paragraphs) for s in list(deck.slides)[1:])
        self.assertEqual(shown, 15)                            # nothing dropped

    def test_content_before_any_heading_is_not_lost(self):
        deck = Presentation(io.BytesIO(doc_slides.render_pptx('Welcome to the plan\n- first point\n')))
        self.assertEqual(deck.slides[0].shapes.title.text, 'Welcome to the plan')


class SheetWriterTests(SimpleTestCase):
    def build(self, text):
        rows = list(csv.reader(io.StringIO(text)))
        return load_workbook(io.BytesIO(doc_sheet.render_xlsx(rows, title='Stock', created=MADE_ON))).active

    def test_values_are_real_and_look_the_same(self):
        sheet = self.build('Item,Qty,Price,Share,Date\nPens,10,"₹1,499.50",12.5%,2026-03-05\n')
        self.assertEqual(sheet['B2'].value, 10)
        self.assertEqual(sheet['C2'].value, 1499.5)
        self.assertEqual(sheet['C2'].number_format, '"₹"#,##0.00')
        self.assertAlmostEqual(sheet['D2'].value, 0.125)
        self.assertEqual(sheet['D2'].number_format, '0.0%')
        self.assertEqual(sheet['E2'].value.date().isoformat(), '2026-03-05')

    def test_identifiers_and_decimal_commas_stay_text(self):
        sheet = self.build('Phone,Code,Ratio\n9876543210,007,"12,34"\n')
        self.assertEqual(sheet['A2'].value, '9876543210')
        self.assertEqual(sheet['B2'].value, '007')
        self.assertEqual(sheet['C2'].value, '12,34')

    def test_formulas_the_model_wrote_stay_live(self):
        sheet = self.build('A,B\n1,2\n=SUM(A2:B2),\n')
        self.assertEqual(sheet['A3'].value, '=SUM(A2:B2)')

    def test_header_style_freeze_filter_and_print_setup(self):
        sheet = self.build('Item,Qty\nPens,10\nTotal,10\n')
        self.assertEqual(sheet['A1'].value, 'Item')
        self.assertEqual(sheet['A1'].fill.fgColor.rgb[-6:].lower(), '059669')
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:B3')
        self.assertEqual(sheet.print_title_rows.replace('$', ''), '1:1')
        self.assertTrue(sheet['A3'].font.bold)                 # the totals row
        self.assertEqual(sheet.title, 'Stock')

    def test_empty_input_still_makes_a_workbook(self):
        self.assertTrue(doc_sheet.render_xlsx([]).startswith(b'PK'))


class TextWriterTests(SimpleTestCase):
    def test_markdown_in_a_txt_becomes_readable_text(self):
        text = doc_text.render_txt(REPORT, brand='Vidhyora', created=MADE_ON)
        self.assertIn('DIGITAL MARKETING PLAN', text)
        self.assertIn('Prepared on 6 October 2026', text)
        self.assertIn('| Meta Ads', text)
        self.assertIn('[NOTE]', text)
        self.assertIn('█', text)                                # the chart as bars
        for marker in ('**', '```', '|---', '](' ):
            self.assertNotIn(marker, text)
        self.assertLessEqual(max(len(line) for line in text.splitlines() if not line.startswith(('+', '|'))), 100)

    def test_plain_text_is_left_exactly_as_written(self):
        self.assertEqual(doc_text.finish_text_file('notes.txt', 'Hello world'), b'Hello world')
        self.assertEqual(doc_text.finish_text_file('list.txt', '- eggs\n- milk\n1. one'), b'- eggs\n- milk\n1. one')

    def test_minified_json_is_indented_but_a_formatted_file_is_not_touched(self):
        self.assertIn(b'\n  "a": 1', doc_text.finish_text_file('d.json', '{"a":1,"b":[1,2]}'))
        formatted = '{\n"a": 1,\n"a": 2\n}\n'
        self.assertEqual(doc_text.finish_text_file('d.json', formatted), formatted.encode())

    def test_csv_with_non_ascii_gets_a_byte_order_mark_for_excel(self):
        self.assertTrue(doc_text.finish_text_file('d.csv', 'a,b\n₹,1').startswith(b'\xef\xbb\xbf'))
        self.assertFalse(doc_text.finish_text_file('d.csv', 'a,b\nx,1').startswith(b'\xef\xbb\xbf'))


class ConversionTests(SimpleTestCase):
    def test_csv_to_pdf_is_a_real_table_not_pipe_text(self):
        data, name, _ = file_convert.convert(b'Item,Qty\nPens,10\n"Books, hardcover",3\n', 'stock.csv', 'pdf')
        text = pdf_text(data)
        self.assertEqual(name, 'stock.pdf')
        self.assertIn('Books, hardcover', text)
        self.assertNotIn(' | ', text)
        self.assertNotIn('Vidhyora', text)                     # a conversion carries no branding

    def test_cells_that_look_like_markdown_are_kept_literally(self):
        data = file_convert.rows_to_pdf_bytes([['Name', 'Note'], ['**bold**', '[x](y)']], branded=False)
        text = pdf_text(data)
        self.assertIn('**bold**', text)
        self.assertIn('[x](y)', text)

    def test_a_huge_table_is_capped_and_says_so(self):
        rows = [['n']] + [[str(i)] for i in range(file_convert.MAX_PDF_TABLE_ROWS + 50)]
        text = pdf_text(file_convert.rows_to_pdf_bytes(rows, branded=False))
        self.assertIn(f'first {file_convert.MAX_PDF_TABLE_ROWS:,} of', text)


class GeneratedFileExtractionTests(SimpleTestCase):
    def test_a_document_may_contain_its_own_code_and_chart_fences(self):
        reply = '```markdown\n# T\n\n```chart\ntype: pie\nA, 1\nB, 2\n```\n\n```python\nprint(1)\n```\n\nEnd\n```'
        content = _extract_ai_generated_file_content(reply, 'report.pdf')
        self.assertIn('```chart', content)
        self.assertIn('print(1)', content)
        self.assertTrue(content.endswith('End'))

    def test_four_backtick_outer_fence(self):
        reply = '````markdown\n# T\n```python\nx\n```\nafter\n````\nbye'
        self.assertEqual(_extract_ai_generated_file_content(reply, 'a.md'), '# T\n```python\nx\n```\nafter')

    def test_text_after_the_closing_fence_is_not_part_of_the_document(self):
        self.assertEqual(_extract_ai_generated_file_content('```markdown\n# T\n\nbody\n```\nHope that helps!', 'a.docx'), '# T\n\nbody')

    def test_a_cut_off_document_keeps_what_was_written(self):
        self.assertEqual(_extract_ai_generated_file_content('```markdown\n# T\n\npartial', 'a.pdf'), '# T\n\npartial')

    def test_code_files_still_take_only_the_first_fence(self):
        self.assertEqual(_extract_ai_generated_file_content('```python\na = 1\n```\n```python\nb = 2\n```', 'x.py'), 'a = 1')

    def test_instructions_ask_for_structure_and_forbid_invention(self):
        for name in ('g.pdf', 'g.docx', 'g.pptx', 'g.xlsx', 'g.txt'):
            text = _ai_generated_file_instruction(name)
            self.assertIn('Never invent statistics', text, msg=name)
            self.assertIn('NEVER write a download link', text, msg=name)
        self.assertIn('## Sources', _ai_generated_file_instruction('g.pdf'))
        self.assertIn('FIRST slide is the title slide', _ai_generated_file_instruction('g.pptx'))
        self.assertIn("'Total'", _ai_generated_file_instruction('g.xlsx'))


class GeneratedFileDownloadTests(TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.user = User.objects.create_user(username='docs-user@example.com', password='test-password-123')
        self.client.force_login(self.user)

    def download(self, name, content):
        stored = AIGeneratedFile.objects.create(user=self.user, file_name=name, content=content)
        return self.client.get(f'/AI/api/files/{stored.token}/download/')

    def test_every_format_downloads_as_a_real_file(self):
        cases = (
            ('report.pdf', REPORT, b'%PDF'), ('report.docx', REPORT, b'PK'),
            ('deck.pptx', '# Title\n- sub\n# Next\n- point\n', b'PK'),
            ('data.xlsx', 'Item,Qty\nPens,10\n', b'PK'),
        )
        for name, content, signature in cases:
            with self.subTest(name=name):
                response = self.download(name, content)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.content.startswith(signature))
                self.assertEqual(response['Cache-Control'], 'private, no-store')

    def test_pdf_download_carries_the_brand_and_the_creation_date(self):
        stored = AIGeneratedFile.objects.create(user=self.user, file_name='report.pdf', content=REPORT)
        text = pdf_text(self.client.get(f'/AI/api/files/{stored.token}/download/').content)
        self.assertIn('Prepared on', text)
        self.assertIn(doc_blocks.date_text(stored.created_at), text)   # the day it was made, in IST

    def test_text_files_are_finished_without_changing_plain_content(self):
        self.assertEqual(self.download('a.txt', 'Hello world').content, b'Hello world')
        self.assertIn('DIGITAL MARKETING PLAN', self.download('r.txt', REPORT).content.decode())
        self.assertTrue(self.download('d.csv', 'a,b\n₹,1').content.startswith(b'\xef\xbb\xbf'))

    def test_a_writer_that_fails_is_a_polite_503_not_a_crash(self):
        with patch('myapp.doc_pdf.render_pdf', side_effect=ValueError('boom')):
            response = self.download('report.pdf', REPORT)
        self.assertEqual(response.status_code, 503)
        self.assertIn('ask for it again', response.json()['detail'])

    def test_helpers_keep_working_without_the_new_options(self):
        self.assertTrue(_ai_pdf_bytes('# T\n\ntext').startswith(b'%PDF'))
        self.assertTrue(_ai_word_document_bytes('# T\n\ntext').startswith(b'PK'))
        self.assertTrue(_ai_excel_bytes('a,b\n1,2').startswith(b'PK'))
        self.assertTrue(_ai_powerpoint_bytes('# T\n- a').startswith(b'PK'))


class CopyCardPromptTests(SimpleTestCase):
    def tag_for(self, text, hint=''):
        reminder = ai_chat.copy_card_reminder(text, hint)
        for tag in ('translation', 'rewrite', 'prompt', 'message'):
            if f'```{tag}' in reminder:
                return tag
        return ''

    def test_requests_that_produce_copyable_text(self):
        for text, expected in (
            ('rephrase this message for me', 'rewrite'), ('make it shorter', 'rewrite'),
            ('translate this in Kannada', 'translation'),
            ('write a prompt for midjourney to create a logo', 'prompt'), ('improve my prompt', 'prompt'),
            ('prompt banao ek poster ke liye', 'prompt'), ('give me a better prompt', 'prompt'),
            ('write me a detailed midjourney prompt for a logo', 'prompt'),
            ('write a WhatsApp message to my customers about the delay', 'message'),
            ('draft an email to my landlord', 'message'),
        ):
            with self.subTest(text=text):
                self.assertEqual(self.tag_for(text), expected)

    def test_ordinary_questions_are_not_turned_into_cards(self):
        for text in (
            'explain photosynthesis in hindi', 'how to write an email to a landlord', 'what is a good prompt for chatgpt',
            'answer in english please', 'write a python function to sort a list', 'hello',
            'write a blog post about prompt engineering', 'make a prompt engineering course outline',
        ):
            with self.subTest(text=text):
                self.assertEqual(self.tag_for(text), '')

    def test_nothing_is_wrapped_during_a_voice_call(self):
        self.assertEqual(self.tag_for('rephrase this', 'This is a live voice call: your reply will be read aloud'), '')

    def test_the_card_instruction_reaches_the_model(self):
        chunk = SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content='ok'))])
        create = Mock(return_value=iter([chunk]))
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
        with patch('myapp.ai_chat._get_client', return_value=client):
            list(ai_chat.stream_chat([{'role': 'user', 'content': 'write a prompt for a bakery logo'}], model_key='quick'))
        system_text = ' '.join(m['content'] for m in create.call_args.kwargs['messages'] if m['role'] == 'system')
        self.assertIn('```prompt', system_text)
        self.assertIn('Copy button', system_text)

    def test_a_plain_question_gets_no_card_instruction(self):
        chunk = SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content='ok'))])
        create = Mock(return_value=iter([chunk]))
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
        with patch('myapp.ai_chat._get_client', return_value=client):
            list(ai_chat.stream_chat([{'role': 'user', 'content': 'what is the capital of France'}], model_key='quick'))
        system_text = ' '.join(m['content'] for m in create.call_args.kwargs['messages'] if m['role'] == 'system')
        self.assertNotIn('```prompt', system_text)
        self.assertNotIn('own Copy button', system_text)


class OfficeXmlOrderTests(SimpleTestCase):
    """Word and PowerPoint refuse — or offer to "repair" — a file whose
    properties sit out of schema order inside an element. The writers add some
    of those by hand, so the order of every property container in the finished
    files is checked here against the schema sequence."""

    WORD_ORDER = {
        'pPr': doc_word._PPR, 'rPr': doc_word._RPR, 'tcPr': doc_word._TCPR,
        'tblPr': doc_word._TBLPR, 'trPr': doc_word._TRPR,
    }
    # DrawingML: names that share a position (any one fill, one underline kind …) share a rank.
    FILL = ('noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill')
    DRAWING_ORDER = {
        'pPr': [('lnSpc',), ('spcBef',), ('spcAft',), ('buClrTx', 'buClr'), ('buSzTx', 'buSzPct', 'buSzPts'),
                ('buFontTx', 'buFont'), ('buNone', 'buAutoNum', 'buChar', 'buBlip'), ('tabLst',), ('defRPr',), ('extLst',)],
        'rPr': [('ln',), FILL, ('effectLst', 'effectDag'), ('highlight',), ('uLnTx', 'uLn'), ('uFillTx', 'uFill'),
                ('latin',), ('ea',), ('cs',), ('sym',), ('hlinkClick',), ('hlinkMouseOver',), ('rtl',), ('extLst',)],
        'tcPr': [('lnL',), ('lnR',), ('lnT',), ('lnB',), ('lnTlToBr',), ('lnBlToTr',), ('cell3D',), FILL,
                 ('headers',), ('extLst',)],
    }

    def assert_in_order(self, root, tag, ranks, where):
        for element in root.iter(tag):
            names = [child.tag.split('}')[1] for child in element if isinstance(child.tag, str)]
            unknown = [name for name in names if name not in ranks]
            self.assertFalse(unknown, f'{where}: unexpected {unknown} inside <{tag.split("}")[1]}>')
            order = [ranks[name] for name in names]
            self.assertEqual(order, sorted(order), f'{where}: {names} out of order inside <{tag.split("}")[1]}>')

    def test_word_properties_are_in_schema_order(self):
        from lxml import etree
        W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        archive = zipfile.ZipFile(io.BytesIO(doc_word.render_docx(REPORT + '\n```python\nx = 1\n```\n', created=MADE_ON)))
        for part in ('word/document.xml', 'word/footer1.xml', 'word/styles.xml'):
            root = etree.fromstring(archive.read(part))
            for local, order in self.WORD_ORDER.items():
                self.assert_in_order(root, W + local, {name: index for index, name in enumerate(order)}, part)

    def test_powerpoint_properties_are_in_schema_order(self):
        from lxml import etree
        A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        deck = SlidesWriterTests.DECK + '\n# Steps\n1. First\n2. Second\n   - sub\nNotes: say hello\n'
        archive = zipfile.ZipFile(io.BytesIO(doc_slides.render_pptx(deck, created=MADE_ON)))
        slides = [name for name in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)]
        self.assertGreaterEqual(len(slides), 5)
        for part in slides:
            root = etree.fromstring(archive.read(part))
            for local, groups in self.DRAWING_ORDER.items():
                ranks = {name: index for index, group in enumerate(groups) for name in group}
                self.assert_in_order(root, A + local, ranks, part)
