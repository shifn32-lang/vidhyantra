"""Designed PowerPoint decks from Markdown-ish text.

A 16:9 deck with a branded title slide, content slides with an accent bar,
text sized to fit (and carried over to a "(cont.)" slide when it cannot),
native tables, native editable charts, speaker notes and slide numbers.

The model writes one slide per heading with its points as bullets (see
views._ai_generated_file_instruction_body). Everything is built on the stock
Title / Title-and-Content layouts, so slides keep their real title and body
placeholders: PowerPoint's outline view, accessibility checker and "reset
slide" all keep working.
"""
import io
import math
import re
from dataclasses import dataclass, field
from datetime import timezone

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from myapp.doc_blocks import (
    Block, CALLOUTS, CHART_COLORS, PALETTE as P, column_aligns, column_weights, date_text,
    format_value, inline_spans, parse_markdown, plain_text, stamp,
)

FONT = 'Calibri'
CODE_FONT = 'Consolas'
INDIC_FONT = 'Nirmala UI'
SLIDE_W, SLIDE_H = 13.333, 7.5            # inches (16:9)
LEFT, BODY_W = 0.85, 11.6
BODY_TOP, BODY_H = 1.85, 4.7
_FONT_SIZES = (26, 24, 22, 20, 18, 16, 14, 12)
_COMFORT_SIZE = 18
_NOTES_RE = re.compile(r'^\s*(?:speaker\s+)?notes?\s*[:：]\s*(.*)$', re.I)


def _rgb(hex6):
    return RGBColor.from_string(hex6.upper())


@dataclass
class _Point:
    text: str
    level: int = 0
    ordered: bool = False
    number: int = 0
    kind: str = 'bullet'      # bullet | plain | quote | code


@dataclass
class _Source:
    title: str
    points: list = field(default_factory=list)
    visuals: list = field(default_factory=list)    # ('table', Block) | ('chart', dict)
    notes: list = field(default_factory=list)


@dataclass
class _Unit:
    """One slide to draw."""
    title: str
    points: list
    visual: tuple = None
    notes: list = field(default_factory=list)
    title_slide: bool = False


# ─────────────────────────────────────────────────────────── blocks -> slides ──

def _sources(blocks):
    """One _Source per heading: its bullets, tables, charts and speaker notes."""
    slides = []

    def add_lines(slide, lines):
        for line in lines:
            note = _NOTES_RE.match(line)
            if note:
                slide.notes.append(plain_text(note.group(1)))
            else:
                slide.points.append(_Point(line))

    for block in blocks:
        if block.kind == 'heading':
            slides.append(_Source(title=plain_text(block.text).replace('\n', ' ')))
            continue
        if block.kind == 'rule':
            continue
        if not slides:
            # Content before any heading still deserves a slide: its first
            # line becomes the title rather than being dropped.
            lines = [line for line in (block.text or '').split('\n') if line.strip()]
            if block.kind in ('para', 'item') and lines:
                slides.append(_Source(title=plain_text(lines[0])[:120]))
                if block.kind == 'para':
                    add_lines(slides[-1], lines[1:])
                continue
            slides.append(_Source(title='Untitled'))
        slide = slides[-1]
        if block.kind == 'item':
            # Continuation lines of one list entry belong to the same bullet.
            text = ' '.join(line.strip() for line in block.text.split('\n') if line.strip())
            note = _NOTES_RE.match(text)
            if note:
                slide.notes.append(plain_text(note.group(1)))
            else:
                slide.points.append(_Point(text, min(block.level, 2), block.ordered, block.number))
        elif block.kind == 'para':
            add_lines(slide, [line for line in block.text.split('\n') if line.strip()])
        elif block.kind == 'quote':
            label = CALLOUTS[block.callout][2] + ': ' if block.callout else ''
            slide.points.append(_Point(label + block.text.replace('\n', ' '), kind='quote'))
        elif block.kind == 'code':
            slide.points.append(_Point(block.text, kind='code'))
        elif block.kind == 'table':
            slide.visuals.append(('table', block))
        elif block.kind == 'chart':
            slide.visuals.append(('chart', block.chart))
    return slides


def _char_width(text):
    """Average glyph width as a fraction of the font size."""
    if any(ord(ch) > 0x2E80 for ch in text):
        return 1.0
    if any(0x0900 <= ord(ch) < 0x0E80 for ch in text):
        return 0.62
    return 0.52


def _points_height(points, size, width_in):
    """Estimated height, in points, of ``points`` at ``size`` pt in ``width_in``."""
    total = 0.0
    for point in points:
        s = size if point.level == 0 else size - 2
        if point.kind == 'code':
            s = max(10, size - 4)
        indent = 0.45 + 0.4 * point.level
        per_line = max(8, (width_in - indent) * 72 / (s * (0.6 if point.kind == 'code' else _char_width(point.text))))
        lines = sum(max(1, math.ceil(len(part) / per_line)) for part in point.text.split('\n'))
        total += lines * s * 1.22 + (9 if point.level == 0 else 4)
    return total


def _fit_size(points, width_in, height_in, floor=0):
    """The largest font size (not below ``floor``) at which ``points`` fit, or None."""
    for size in _FONT_SIZES:
        if size < floor:
            break
        if _points_height(points, size, width_in) <= height_in * 72:
            return size
    return None


def _paginate_points(points, width_in, height_in):
    """Split ``points`` into groups that each fit at a size people can read
    from the back of a room (at least _COMFORT_SIZE pt). A single point too long
    for that is kept whole and simply set smaller."""
    if not points or _fit_size(points, width_in, height_in, _COMFORT_SIZE) is not None:
        return [points] if points else [[]]
    pages, page = [], []
    for point in points:
        trial = page + [point]
        if page and _fit_size(trial, width_in, height_in, _COMFORT_SIZE) is None:
            pages.append(page)
            page = [point]
        else:
            page = trial
    if page:
        pages.append(page)
    return pages


def _table_rows_per_slide(block, available_in):
    columns = len(block.header)
    size = 16 if columns <= 4 else 14 if columns <= 6 else 12
    weights = column_weights(block)
    widths = [BODY_W * w / (sum(weights) or 1) for w in weights]

    def row_height(cells):
        lines = 1
        for text, width in zip(cells, widths):
            per_line = max(4, (width - 0.3) * 72 / (size * _char_width(text)))
            lines = max(lines, math.ceil(len(plain_text(text)) / per_line))
        return (lines * size * 1.25 + 14) / 72

    pages, page, used = [], [], row_height(block.header)
    for row in block.rows:
        height = row_height(row)
        if page and used + height > available_in:
            pages.append(page)
            page, used = [], row_height(block.header)
        page.append(row)
        used += height
    pages.append(page)
    return pages, size, widths


def _units(sources):
    """Slides to draw. A heading with only a line or two of text first in the
    deck becomes the title slide; anything too long for one slide continues on
    a "(cont.)" slide."""
    units = []
    for index, source in enumerate(sources):
        point_pages = _paginate_points(source.points, BODY_W, BODY_H)
        visuals = list(source.visuals)
        combine = False
        if len(visuals) == 1 and len(point_pages) == 1 and 0 < len(point_pages[0]) <= 4:
            if visuals[0][0] == 'chart':
                combine = _fit_size(point_pages[0], 4.4, BODY_H) is not None
            else:
                combine = len(visuals[0][1].rows) <= 5 and _fit_size(point_pages[0], BODY_W, 1.35) is not None
        made = []
        if combine:
            made.append(_Unit(source.title, point_pages[0], visuals[0]))
        else:
            for page in point_pages:
                if page or not visuals:
                    made.append(_Unit(source.title, page))
            for kind, payload in visuals:
                if kind == 'table':
                    row_pages, _size, _widths = _table_rows_per_slide(payload, BODY_H)
                    for rows in row_pages:
                        part = Block('table', header=payload.header, rows=rows, aligns=payload.aligns)
                        made.append(_Unit(source.title, [], ('table', part)))
                else:
                    made.append(_Unit(source.title, [], (kind, payload)))
        for number, unit in enumerate(made):
            if number > 0:
                unit.title = f'{source.title} (cont.)'
        made[0].notes = source.notes
        if (
            index == 0 and len(made) == 1 and made[0].visual is None and len(made[0].points) <= 2
            and all(len(p.text) <= 140 and p.kind == 'bullet' for p in made[0].points)
        ):
            made[0].title_slide = True
        units.extend(made)
    return units


# ───────────────────────────────────────────────────────────── drawing bits ──

def _add_cs_font(run):
    """Complex-script font (Hindi, Tamil …) next to the Latin one."""
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:cs',):
        for existing in rpr.findall(qn(tag)):
            rpr.remove(existing)
    cs = rpr.makeelement(qn('a:cs'), {'typeface': INDIC_FONT})
    after = [qn('a:sym'), qn('a:hlinkClick'), qn('a:hlinkMouseOver'), qn('a:rtl'), qn('a:extLst')]
    for index, child in enumerate(rpr):
        if child.tag in after:
            rpr.insert(index, cs)
            return
    rpr.append(cs)


def _style_run(run, size, color, bold=False, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold or None
    run.font.italic = italic or None
    run.font.color.rgb = _rgb(color)
    _add_cs_font(run)


def _add_spans(paragraph, text, size, color, *, bold=False, italic=False):
    for span in inline_spans(text):
        run = paragraph.add_run()
        run.text = span.text.replace('\n', ' ')
        if span.code:
            _style_run(run, max(10, size - 2), P['accent_dark'], bold, italic, CODE_FONT)
        else:
            _style_run(run, size, P['accent_dark'] if span.url else color, bold or span.bold, italic or span.italic)
        if span.strike:
            run._r.get_or_add_rPr().set('strike', 'sngStrike')
        if span.url:
            run.hyperlink.address = span.url
            run.font.underline = True


def _bullet(paragraph, point, first_of_run):
    """Bullet character, colour and hanging indent for one paragraph. The order
    of the children — colour, font, then the character — is the schema's."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ('a:buClr', 'a:buSzPct', 'a:buFont', 'a:buNone', 'a:buAutoNum', 'a:buChar'):
        for existing in pPr.findall(qn(tag)):
            pPr.remove(existing)
    if point.kind in ('plain', 'quote', 'code'):
        pPr.set('marL', '0')
        pPr.set('indent', '0')
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))
        return
    left = 342900 + 342900 * point.level
    pPr.set('marL', str(left))
    pPr.set('indent', str(-285750))
    pPr.append(_solid(pPr, 'a:buClr', P['accent']))
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
    if point.ordered:
        attrs = {'type': 'arabicPeriod'}
        if first_of_run and point.number > 1:
            attrs['startAt'] = str(point.number)
        pPr.append(pPr.makeelement(qn('a:buAutoNum'), attrs))
    else:
        pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '■' if point.level == 0 else '–'}))


def _solid(parent, tag, hex6):
    element = parent.makeelement(qn(tag), {})
    element.append(element.makeelement(qn('a:srgbClr'), {'val': hex6.upper()}))
    return element


def _rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _textbox(slide, x, y, w, h, text, size, color, *, bold=False, align=PP_ALIGN.LEFT, spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size, color, bold)
    if spacing:
        run._r.get_or_add_rPr().set('spc', str(int(spacing * 100)))
    return box


def _slide_number(slide, number, x, y, w, h, color):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.margin_right = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    field_el = paragraph._p.makeelement(qn('a:fld'), {'id': '{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}', 'type': 'slidenum'})
    rpr = field_el.makeelement(qn('a:rPr'), {'lang': 'en-US', 'sz': '1000'})
    rpr.append(_solid(rpr, 'a:solidFill', color))
    text = field_el.makeelement(qn('a:t'), {})
    text.text = str(number)
    field_el.append(rpr)
    field_el.append(text)
    paragraph._p.append(field_el)


def _title_size(text, large):
    n = len(text)
    if large:
        return 44 if n <= 28 else 38 if n <= 48 else 32 if n <= 80 else 26
    return 32 if n <= 42 else 28 if n <= 66 else 24 if n <= 100 else 20


def _set_geometry(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def _fill_title(shape, text, size, color, *, anchor=MSO_ANCHOR.MIDDLE):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text[:250]
    _style_run(run, size, color, bold=True)


def _fill_body(shape, points, size, x, y, w, h):
    _set_geometry(shape, x, y, w, h)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = 0
    previous = None
    for index, point in enumerate(points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        text_size = size if point.level == 0 else size - 2
        if point.kind == 'code':
            text_size = max(10, size - 4)
            run = paragraph.add_run()
            run.text = point.text
            _style_run(run, text_size, P['ink'], name=CODE_FONT)
        elif point.kind == 'quote':
            _add_spans(paragraph, point.text, text_size, P['muted'], italic=True)
        else:
            _add_spans(paragraph, point.text, text_size, P['body'])
        paragraph.space_after = Pt(9 if point.level == 0 else 4)
        paragraph.line_spacing = 1.05
        paragraph.alignment = PP_ALIGN.LEFT
        first = not (previous and previous.ordered and point.ordered and previous.level == point.level)
        _bullet(paragraph, point, first)
        previous = point


def _remove(shape):
    shape._element.getparent().remove(shape._element)


def _add_table(slide, block, x, y, w, size, widths):
    columns = len(block.header)
    aligns = column_aligns(block)
    align = {'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}
    row_h = (size * 1.25 + 14) / 72
    frame = slide.shapes.add_table(1 + len(block.rows), columns, Inches(x), Inches(y), Inches(w), Inches(row_h * (1 + len(block.rows))))
    table = frame.table
    table.first_row = True
    table.horz_banding = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)

    def style_cell(cell, text, header, shade, bold):
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = align[aligns[column]]
        _add_spans(paragraph, text, size, P['white'] if header else P['body'], bold=header or bold)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(P['accent'] if header else (shade or P['white']))
        # Cell edges: only a thin line under each row. The edges must come
        # before the fill inside tcPr, so they are inserted at the front.
        properties = cell._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            for existing in properties.findall(qn(tag)):
                properties.remove(existing)
        edges = []
        for tag, color in (('a:lnL', None), ('a:lnR', None), ('a:lnT', None), ('a:lnB', P['rule'])):
            line = properties.makeelement(qn(tag), {'w': '9525' if color else '0'})
            line.append(_solid(line, 'a:solidFill', color) if color else line.makeelement(qn('a:noFill'), {}))
            edges.append(line)
        for position, line in enumerate(edges):
            properties.insert(position, line)

    for column, text in enumerate(block.header):
        style_cell(table.cell(0, column), text, True, None, False)
    table.rows[0].height = Inches(row_h)
    for number, cells in enumerate(block.rows, start=1):
        total = bool(cells) and plain_text(cells[0]).strip().lower().startswith(('total', 'grand total'))
        shade = P['soft'] if total else (P['zebra'] if number % 2 == 0 else None)
        for column, text in enumerate(cells):
            style_cell(table.cell(number, column), text, False, shade, total)
        table.rows[number].height = Inches(row_h)
    return frame


def _add_chart(slide, spec, x, y, w, h):
    data = CategoryChartData()
    data.categories = [label or ' ' for label in spec['labels']]
    for series in spec['series']:
        data.add_series(series['name'], series['values'])
    kind = {
        'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'hbar': XL_CHART_TYPE.BAR_CLUSTERED,
        'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE, 'donut': XL_CHART_TYPE.DOUGHNUT,
    }[spec['type']]
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.font.size = Pt(13)
    chart.font.name = FONT
    chart.font.color.rgb = _rgb(P['body'])
    title = ' — '.join(part for part in (spec['title'], spec['unit']) if part)
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = _rgb(P['ink'])
    circular = spec['type'] in ('pie', 'donut')
    chart.has_legend = circular or len(spec['series']) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if circular else XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    plot = chart.plots[0]
    few = len(spec['labels']) * len(spec['series']) <= 16
    if circular:
        plot.vary_by_categories = True
        for index, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(CHART_COLORS[index % len(CHART_COLORS)])
            point.format.line.color.rgb = _rgb(P['white'])
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_percentage = True
        labels.show_value = False
        labels.number_format = '0%'
        labels.number_format_is_linked = False
        labels.font.size = Pt(13)
        labels.font.bold = True
        labels.font.color.rgb = _rgb(P['white'])
    else:
        for index, series in enumerate(plot.series):
            color = _rgb(CHART_COLORS[index % len(CHART_COLORS)])
            if spec['type'] == 'line':
                series.format.line.color.rgb = color
                series.format.line.width = Pt(3)
                series.smooth = False
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
        if few:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = '#,##0.##'
            labels.number_format_is_linked = False
            labels.font.size = Pt(12)
            labels.font.bold = True
            if spec['type'] != 'line':
                labels.position = XL_LABEL_POSITION.OUTSIDE_END
            else:
                labels.position = XL_LABEL_POSITION.ABOVE
        axis = chart.value_axis
        axis.has_major_gridlines = True
        axis.major_gridlines.format.line.color.rgb = _rgb('e5e7eb')
        axis.format.line.fill.background()
        axis.tick_labels.font.size = Pt(12)
        axis.tick_labels.font.color.rgb = _rgb(P['muted'])
        chart.category_axis.tick_labels.font.size = Pt(12)
        chart.category_axis.format.line.color.rgb = _rgb(P['faint'])
        if spec['type'] in ('column', 'hbar'):
            plot.gap_width = 70
    return frame


# ─────────────────────────────────────────────────────────────── rendering ──

def _title_slide(prs, unit, brand, created, deck_title):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(P['deep'])
    title, subtitle = slide.shapes.title, slide.placeholders[1]
    _set_geometry(title, 0.9, 2.0, 11.5, 2.2)
    _fill_title(title, unit.title, _title_size(unit.title, True), P['white'], anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, 0.9, 4.4, 1.4, 0.08, P['mint'])
    _set_geometry(subtitle, 0.9, 4.7, 11.0, 1.6)
    frame = subtitle.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = 0
    for index, point in enumerate(unit.points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(6)
        _add_spans(paragraph, point.text, 22, P['mint'])
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set('marL', '0')
        pPr.set('indent', '0')
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))
    if brand:
        _textbox(slide, 0.9, 0.85, 8, 0.4, f'{brand.upper()} AI', 12, P['mint'], bold=True, spacing=2)
    _textbox(slide, 0.9, 6.75, 8, 0.4, f'Prepared on {date_text(created)}', 12, P['mint'])
    return slide


def _content_slide(prs, unit, number, brand):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _rect(slide, 0, 0, 0.25, SLIDE_H, P['accent'])
    title, body = slide.shapes.title, slide.placeholders[1]
    _set_geometry(title, LEFT, 0.4, BODY_W, 1.05)
    _fill_title(title, unit.title, _title_size(unit.title, False), P['ink'])
    _rect(slide, LEFT, 1.5, 1.1, 0.07, P['accent'])
    if brand:
        _textbox(slide, LEFT, 6.95, 6, 0.3, f'{brand} AI', 10, P['faint'])
    _slide_number(slide, number, 11.4, 6.95, 1.35, 0.3, P['faint'])

    visual = unit.visual
    points = unit.points
    if visual is None:
        if points:
            size = _fit_size(points, BODY_W, BODY_H) or _FONT_SIZES[-1]
            _fill_body(body, points, size, LEFT, BODY_TOP, BODY_W, BODY_H)
        else:
            _remove(body)
    elif visual[0] == 'chart':
        if points:
            size = min(_fit_size(points, 4.4, BODY_H) or 14, 22)
            _fill_body(body, points, size, LEFT, BODY_TOP, 4.4, BODY_H)
            _add_chart(slide, visual[1], 5.5, BODY_TOP - 0.1, 6.95, BODY_H + 0.1)
        else:
            _remove(body)
            _add_chart(slide, visual[1], LEFT, BODY_TOP - 0.1, BODY_W, BODY_H + 0.1)
    else:   # table
        block = visual[1]
        _pages, size, widths = _table_rows_per_slide(block, BODY_H)
        top = BODY_TOP
        if points:
            height = 1.35
            size_text = min(_fit_size(points, BODY_W, height) or 14, 20)
            _fill_body(body, points, size_text, LEFT, BODY_TOP, BODY_W, height)
            top = BODY_TOP + height + 0.15
        else:
            _remove(body)
        _add_table(slide, block, LEFT, top, BODY_W, size, widths)
    return slide


def render_pptx(content, *, title='', brand='', created=None, blocks=None):
    """Markdown-ish text -> a .pptx as bytes. One slide per heading."""
    blocks = parse_markdown(content) if blocks is None else blocks
    units = _units(_sources(blocks))
    if not units:
        units = [_Unit(title or 'Untitled', [])]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)

    deck_title = units[0].title
    for number, unit in enumerate(units, start=1):
        if unit.title_slide:
            slide = _title_slide(prs, unit, brand, created, deck_title)
        else:
            slide = _content_slide(prs, unit, number, brand)
        if unit.notes:
            slide.notes_slide.notes_text_frame.text = '\n'.join(unit.notes)

    moment = stamp(created)
    props = prs.core_properties
    props.title = deck_title
    props.author = brand
    props.last_modified_by = brand
    props.comments = f'Created with {brand} AI' if brand else ''
    props.created = props.modified = moment.astimezone(timezone.utc).replace(tzinfo=None)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
